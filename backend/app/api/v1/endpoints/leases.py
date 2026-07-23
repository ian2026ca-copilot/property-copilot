import os
import uuid
import pathlib
from datetime import date
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.core.database import get_db
from app.api.deps import get_current_user, require_min_role
from app.models.user import User, OrganizationMember, UserRole, TenantDocument
from app.models.profiles import TenantProfile
from app.models.lease import Lease, LeaseStatus, LeaseType
from app.models.payment import Payment
from app.models.lease_template import LeaseTemplate
from app.models.property import Unit, Property, UnitStatus
from app.schemas.lease import LeaseCreate, LeaseUpdate, LeaseRenew, LeaseOut, TenantOut, TenantDocumentOut
from app.api.v1.endpoints.payments import generate_monthly_payments

router = APIRouter(prefix="/leases", tags=["leases"])

UPLOAD_DIR = pathlib.Path("/app/uploads")
LEASE_DOC_DIR = UPLOAD_DIR / "leases"
MAX_SIZE_MB = 50
ALLOWED_DOC_TYPES = {
    "application/pdf",
    "image/jpeg",
    "image/png",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def _uploads_url(filename: str) -> str:
    base = os.getenv("FRONTEND_URL", "http://localhost:3001").replace(":3001", ":8002").replace(":3000", ":8002")
    return f"{base}/uploads/{filename}"


def _doc_url(path: str) -> str:
    """Convert a stored document_path to a public URL."""
    base = os.getenv("FRONTEND_URL", "http://localhost:3001").replace(":3001", ":8002").replace(":3000", ":8002")
    return f"{base}/uploads/{path}"


def _compute_status_from_dates(start_date: date, end_date: date) -> LeaseStatus:
    today = date.today()
    if start_date > today:
        return LeaseStatus.PENDING
    if end_date < today:
        return LeaseStatus.EXPIRED
    return LeaseStatus.ACTIVE


def _effective_status(lease: Lease) -> LeaseStatus:
    """Return stored status directly â€” status is computed on create/date-update."""
    return lease.status


async def _upsert_tenant_profile(db: AsyncSession, user_id, **fields) -> TenantProfile:
    """Create or update the given user's TenantProfile with any non-None fields provided."""
    result = await db.execute(select(TenantProfile).where(TenantProfile.user_id == user_id))
    profile = result.scalar_one_or_none()
    if not profile:
        profile = TenantProfile(user_id=user_id)
        db.add(profile)
    for field, value in fields.items():
        if value is not None:
            setattr(profile, field, value)
    return profile


def _tenant_to_out(user: User, docs: list[TenantDocument] | None = None) -> TenantOut:
    profile = user.tenant_profile
    return TenantOut(
        id=user.id,
        first_name=user.first_name,
        last_name=user.last_name,
        full_name=user.display_name,
        email=user.email,
        phone=user.phone or "",
        date_of_birth=profile.date_of_birth if profile else None,
        avatar_url=_uploads_url(user.avatar_filename) if user.avatar_filename else None,
        documents=[
            TenantDocumentOut(
                id=d.id,
                doc_type=d.doc_type,
                filename=d.filename,
                original_name=d.original_name,
                url=_uploads_url(d.filename),
            )
            for d in (docs or [])
        ],
    )


def _lease_to_out(lease: Lease) -> LeaseOut:
    docs = list(lease.tenant.tenant_documents) if lease.tenant and hasattr(lease.tenant, "tenant_documents") else []
    doc_url = _doc_url(lease.document_path) if lease.document_path else None
    co_tenants = [_tenant_to_out(u) for u in (lease.co_tenants or [])]
    return LeaseOut(
        id=lease.id,
        unit_id=lease.unit_id,
        tenant_user_id=lease.tenant_user_id,
        start_date=lease.start_date,
        end_date=lease.end_date,
        monthly_rent=lease.monthly_rent,
        security_deposit=lease.security_deposit,
        status=_effective_status(lease),
        lease_type=lease.lease_type,
        document_url=doc_url,
        landlord_name=lease.landlord_name,
        landlord_email=lease.landlord_email,
        docusign_envelope_id=lease.docusign_envelope_id,
        signature_status=lease.signature_status,
        notes=lease.notes,
        tenant=_tenant_to_out(lease.tenant, docs) if lease.tenant else None,
        co_tenants=co_tenants,
        unit_number=lease.unit.unit_number if lease.unit else None,
        property_name=lease.unit.property.name if lease.unit and lease.unit.property else None,
    )


async def _get_lease(lease_id: str, org_id: uuid.UUID, db: AsyncSession) -> Lease:
    result = await db.execute(
        select(Lease)
        .where(Lease.id == lease_id, Lease.organization_id == org_id)
        .options(
            selectinload(Lease.tenant).selectinload(User.tenant_documents),
            selectinload(Lease.tenant).selectinload(User.tenant_profile),
            selectinload(Lease.co_tenants).selectinload(User.tenant_profile),
            selectinload(Lease.unit).selectinload(Unit.property),
        )
    )
    lease = result.scalar_one_or_none()
    if not lease:
        raise HTTPException(status_code=404, detail="Lease not found")
    return lease




# â”€â”€ Lease Templates â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

TEMPLATE_DIR = UPLOAD_DIR / "lease_templates"
ALLOWED_TEMPLATE_TYPES = {
    "application/pdf",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "image/jpeg",
    "image/png",
}


class LeaseTemplateOut(BaseModel):
    id: str
    name: str
    original_name: str
    description: Optional[str]
    url: str
    created_at: str

    class Config:
        from_attributes = True


# â”€â”€ List â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

@router.get("", response_model=list[LeaseOut])
async def list_leases(
    current: tuple[User, OrganizationMember] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    result = await db.execute(
        select(Lease)
        .where(Lease.organization_id == member.organization_id)
        .options(
            selectinload(Lease.tenant).selectinload(User.tenant_documents),
            selectinload(Lease.tenant).selectinload(User.tenant_profile),
            selectinload(Lease.co_tenants).selectinload(User.tenant_profile),
            selectinload(Lease.unit).selectinload(Unit.property),
        )
        .order_by(Lease.created_at.desc())
    )
    return [_lease_to_out(l) for l in result.scalars().all()]


# â”€â”€ Create â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

@router.post("", response_model=LeaseOut, status_code=status.HTTP_201_CREATED)
async def create_lease(
    body: LeaseCreate,
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    _, member = current

    # Verify tenant exists
    tenant_result = await db.execute(select(User).where(User.id == body.tenant_user_id, User.is_active == True))
    if not tenant_result.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Tenant not found")

    # Verify unit belongs to org
    unit_result = await db.execute(
        select(Unit).join(Property).where(
            Unit.id == body.unit_id,
            Property.organization_id == member.organization_id,
            Property.is_active == True,
        )
    )
    unit = unit_result.scalar_one_or_none()
    if not unit:
        raise HTTPException(status_code=404, detail="Unit not found")

    lease = Lease(
        organization_id=member.organization_id,
        unit_id=body.unit_id,
        tenant_user_id=body.tenant_user_id,
        start_date=body.start_date,
        end_date=body.end_date,
        monthly_rent=body.monthly_rent,
        security_deposit=body.security_deposit,
        lease_type=body.lease_type,
        landlord_name=body.landlord_name or None,
        notes=body.notes,
        status=_compute_status_from_dates(body.start_date, body.end_date),
    )
    db.add(lease)
    await db.flush()

    # Mark unit occupied if lease is active
    if lease.status == LeaseStatus.ACTIVE:
        unit.status = UnitStatus.OCCUPIED

    # Generate monthly rent payments
    monthly_payments = generate_monthly_payments(lease, member.organization_id)
    for p in monthly_payments:
        db.add(p)

    await db.flush()

    # Add co-tenants
    if body.co_tenant_ids:
        from app.models.lease import lease_co_tenants
        from sqlalchemy import insert as sa_insert
        valid_co_ids = [cid for cid in body.co_tenant_ids if cid != body.tenant_user_id]
        if valid_co_ids:
            await db.execute(
                sa_insert(lease_co_tenants),
                [{"lease_id": lease.id, "user_id": cid} for cid in valid_co_ids],
            )

    await db.commit()
    return _lease_to_out(await _get_lease(str(lease.id), member.organization_id, db))


# â”€â”€ Read one â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
def _template_url(file_path: str) -> str:
    base = os.environ.get("BACKEND_URL", "http://localhost:8000")
    return f"{base}/uploads/{file_path}"


@router.get("/templates", response_model=list[LeaseTemplateOut])
async def list_lease_templates(
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    result = await db.execute(
        select(LeaseTemplate)
        .where(LeaseTemplate.organization_id == member.organization_id)
        .order_by(LeaseTemplate.created_at.desc())
    )
    templates = result.scalars().all()
    return [
        LeaseTemplateOut(
            id=str(t.id),
            name=t.name,
            original_name=t.original_name,
            description=t.description,
            url=_template_url(t.file_path),
            created_at=t.created_at.isoformat(),
        )
        for t in templates
    ]


@router.post("/templates", response_model=LeaseTemplateOut, status_code=201)
async def upload_lease_template(
    file: UploadFile = File(...),
    name: str = Form(...),
    description: str = Form(""),
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    if file.content_type not in ALLOWED_TEMPLATE_TYPES:
        raise HTTPException(status_code=400, detail="Only PDF, Word, JPEG, or PNG files accepted")
    data = await file.read()
    if len(data) > MAX_SIZE_MB * 1024 * 1024:
        raise HTTPException(status_code=400, detail=f"File exceeds {MAX_SIZE_MB} MB limit")

    ext = pathlib.Path(file.filename or "template").suffix or ".pdf"
    filename = f"lease_templates/template_{uuid.uuid4()}{ext}"
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    (UPLOAD_DIR / filename).write_bytes(data)

    tmpl = LeaseTemplate(
        organization_id=member.organization_id,
        name=name,
        original_name=file.filename or name,
        file_path=filename,
        description=description or None,
    )
    db.add(tmpl)
    await db.commit()
    await db.refresh(tmpl)
    return LeaseTemplateOut(
        id=str(tmpl.id),
        name=tmpl.name,
        original_name=tmpl.original_name,
        description=tmpl.description,
        url=_template_url(tmpl.file_path),
        created_at=tmpl.created_at.isoformat(),
    )


@router.post("/templates/ai-generate", response_model=LeaseTemplateOut, status_code=201)
async def ai_generate_lease_template(
    province: str = Form(...),
    lease_type: str = Form("Fixed-term"),
    property_type: str = Form("Residential Apartment"),
    bedrooms: str = Form(""),
    notes: str = Form(""),
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    """Use Gemini to generate a lease agreement template as a Word document."""
    import io
    import google.generativeai as genai
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    api_key = os.environ.get("GEMINI_API_KEY", "")
    if not api_key:
        raise HTTPException(status_code=503, detail="Gemini API key not configured")

    _, member = current

    bedrooms_line = f"Bedrooms: {bedrooms}\n" if bedrooms else ""
    notes_line = f"Additional requirements: {notes}\n" if notes else ""

    prompt = (
        f"Generate a comprehensive residential lease agreement for {province}, Canada.\n\n"
        f"Lease type: {lease_type}\n"
        f"Property type: {property_type}\n"
        f"{bedrooms_line}"
        f"{notes_line}\n"
        "Use these placeholders throughout: [LANDLORD NAME], [TENANT NAME], [PROPERTY ADDRESS], "
        "[MONTHLY RENT], [SECURITY DEPOSIT], [START DATE], [END DATE].\n\n"
        "Structure with these numbered sections (use ALL CAPS for section titles):\n"
        "1. PARTIES\n2. PREMISES\n3. TERM\n4. RENT\n5. SECURITY DEPOSIT\n"
        "6. UTILITIES AND SERVICES\n7. MAINTENANCE AND REPAIRS\n8. ENTRY BY LANDLORD\n"
        "9. PETS POLICY\n10. SMOKING POLICY\n11. ALTERATIONS\n12. SUBLETTING\n"
        "13. TERMINATION\n14. DEFAULT\n15. GOVERNING LAW\n16. SIGNATURES\n\n"
        f"Make it legally appropriate for {province} residential tenancy law.\n"
        "Return plain text only â€” no markdown, no backticks."
    )

    try:
        genai.configure(api_key=api_key)
        model = genai.GenerativeModel("gemini-2.5-flash")
        response = model.generate_content(prompt)
        content = response.text or ""
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Gemini error: {str(e)[:200]}")

    # Build Word document
    doc = Document()

    # Title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("RESIDENTIAL LEASE AGREEMENT")
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph()

    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph()
            continue
        # Section header detection: ALL CAPS line or starts with digit+dot
        is_header = stripped.isupper() and len(stripped) > 3
        if not is_header and len(stripped) > 2:
            parts = stripped.split(".", 1)
            is_header = parts[0].strip().isdigit()
        if is_header:
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(11)
        else:
            doc.add_paragraph(stripped)

    # Signature block
    doc.add_paragraph()
    sig = doc.add_paragraph()
    sig.add_run("LANDLORD: ________________________________    Date: ____________").bold = False
    doc.add_paragraph()
    sig2 = doc.add_paragraph()
    sig2.add_run("TENANT:   ________________________________    Date: ____________").bold = False

    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    template_name = f"AI â€” {property_type} ({province})"
    filename = f"lease_templates/ai_lease_{uuid.uuid4()}.docx"
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    (UPLOAD_DIR / filename).write_bytes(data)

    tmpl = LeaseTemplate(
        organization_id=member.organization_id,
        name=template_name,
        original_name=f"{template_name}.docx",
        file_path=filename,
        description=f"{lease_type} | {property_type} | {province}" + (f" | {notes}" if notes else ""),
    )
    db.add(tmpl)
    await db.commit()
    await db.refresh(tmpl)
    return LeaseTemplateOut(
        id=str(tmpl.id),
        name=tmpl.name,
        original_name=tmpl.original_name,
        description=tmpl.description,
        url=_template_url(tmpl.file_path),
        created_at=tmpl.created_at.isoformat(),
    )


@router.delete("/templates/{template_id}", status_code=204)
async def delete_lease_template(
    template_id: str,
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    result = await db.execute(
        select(LeaseTemplate).where(
            LeaseTemplate.id == template_id,
            LeaseTemplate.organization_id == member.organization_id,
        )
    )
    tmpl = result.scalar_one_or_none()
    if not tmpl:
        raise HTTPException(status_code=404, detail="Template not found")
    try:
        (UPLOAD_DIR / tmpl.file_path).unlink(missing_ok=True)
    except OSError:
        pass
    await db.delete(tmpl)
    await db.commit()


@router.get("/{lease_id}", response_model=LeaseOut)
async def get_lease(
    lease_id: str,
    current: tuple[User, OrganizationMember] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    return _lease_to_out(await _get_lease(lease_id, member.organization_id, db))


# â”€â”€ Update â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

@router.put("/{lease_id}", response_model=LeaseOut)
async def update_lease(
    lease_id: str,
    body: LeaseUpdate,
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    lease = await _get_lease(lease_id, member.organization_id, db)

    data = body.model_dump(exclude_none=True)
    tenant_fields = {"first_name", "last_name", "phone"}
    date_of_birth = data.pop("date_of_birth", None)

    if lease.tenant:
        for f in tenant_fields:
            if f in data:
                setattr(lease.tenant, f, data.pop(f))
        if lease.tenant.first_name or lease.tenant.last_name:
            lease.tenant.full_name = " ".join(filter(None, [lease.tenant.first_name, lease.tenant.last_name]))
        if date_of_birth:
            await _upsert_tenant_profile(db, lease.tenant.id, date_of_birth=date_of_birth)
    else:
        for f in tenant_fields:
            data.pop(f, None)

    # If no explicit status override, recompute from dates when dates change
    date_changed = "start_date" in data or "end_date" in data
    explicit_status = "status" in data

    for field, value in data.items():
        setattr(lease, field, value)

    if not explicit_status and date_changed and lease.status != LeaseStatus.TERMINATED:
        lease.status = _compute_status_from_dates(lease.start_date, lease.end_date)

    # Sync unit status based on effective lease status
    if lease.unit:
        if lease.status == LeaseStatus.ACTIVE:
            lease.unit.status = UnitStatus.OCCUPIED
        elif lease.status == LeaseStatus.TERMINATED:
            lease.unit.status = UnitStatus.VACANT

    await db.commit()
    return _lease_to_out(await _get_lease(lease_id, member.organization_id, db))


# â”€â”€ Terminate â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

@router.delete("/{lease_id}", status_code=status.HTTP_204_NO_CONTENT)
async def terminate_lease(
    lease_id: str,
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    result = await db.execute(
        select(Lease).where(Lease.id == lease_id, Lease.organization_id == member.organization_id)
    )
    lease = result.scalar_one_or_none()
    if not lease:
        raise HTTPException(status_code=404, detail="Lease not found")

    lease.status = LeaseStatus.TERMINATED

    # Restore unit to vacant if no other active lease
    other = await db.execute(
        select(Lease).where(
            Lease.unit_id == lease.unit_id,
            Lease.id != lease.id,
            Lease.status.in_([LeaseStatus.ACTIVE, LeaseStatus.PENDING]),
        )
    )
    if not other.scalar_one_or_none():
        unit_res = await db.execute(select(Unit).where(Unit.id == lease.unit_id))
        unit = unit_res.scalar_one_or_none()
        if unit:
            unit.status = UnitStatus.VACANT

    await db.commit()


@router.delete("/{lease_id}/permanent", status_code=status.HTTP_204_NO_CONTENT)
async def delete_lease_permanent(
    lease_id: str,
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    result = await db.execute(
        select(Lease).where(Lease.id == lease_id, Lease.organization_id == member.organization_id)
    )
    lease = result.scalar_one_or_none()
    if not lease:
        raise HTTPException(status_code=404, detail="Lease not found")

    # Restore unit to vacant if no other active lease remains
    other = await db.execute(
        select(Lease).where(
            Lease.unit_id == lease.unit_id,
            Lease.id != lease.id,
            Lease.status.in_([LeaseStatus.ACTIVE, LeaseStatus.PENDING]),
        )
    )
    if not other.scalars().first():
        unit_res = await db.execute(select(Unit).where(Unit.id == lease.unit_id))
        unit = unit_res.scalar_one_or_none()
        if unit:
            unit.status = UnitStatus.VACANT

    doc_path = lease.document_path
    lease_id_val = lease.id

    # Use raw SQL deletes to avoid ORM cascade/lazy-load issues in async context
    from sqlalchemy import delete as sa_delete
    await db.execute(sa_delete(Payment).where(Payment.lease_id == lease_id_val))
    await db.execute(sa_delete(Lease).where(Lease.id == lease_id_val))
    await db.commit()

    if doc_path:
        try:
            (UPLOAD_DIR / doc_path).unlink(missing_ok=True)
        except OSError:
            pass



# â”€â”€ Renew â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

@router.post("/{lease_id}/renew", response_model=LeaseOut, status_code=status.HTTP_201_CREATED)
async def renew_lease(
    lease_id: str,
    body: LeaseRenew,
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    old = await _get_lease(lease_id, member.organization_id, db)

    new_unit_id = body.unit_id if body.unit_id is not None else old.unit_id
    new_tenant_id = body.tenant_user_id if body.tenant_user_id is not None else old.tenant_user_id

    if body.tenant_user_id is not None:
        tenant_result = await db.execute(select(User).where(User.id == body.tenant_user_id, User.is_active == True))
        if not tenant_result.scalar_one_or_none():
            raise HTTPException(status_code=404, detail="Tenant not found")

    if body.unit_id is not None:
        unit_result = await db.execute(
            select(Unit).join(Property).where(
                Unit.id == body.unit_id,
                Property.organization_id == member.organization_id,
                Property.is_active == True,
            )
        )
        if not unit_result.scalar_one_or_none():
            raise HTTPException(status_code=404, detail="Unit not found")

    old.status = LeaseStatus.TERMINATED

    new_lease = Lease(
        organization_id=member.organization_id,
        unit_id=new_unit_id,
        tenant_user_id=new_tenant_id,
        start_date=body.start_date,
        end_date=body.end_date,
        monthly_rent=body.monthly_rent if body.monthly_rent is not None else old.monthly_rent,
        security_deposit=body.security_deposit if body.security_deposit is not None else old.security_deposit,
        lease_type=body.lease_type if body.lease_type is not None else old.lease_type,
        landlord_name=body.landlord_name if body.landlord_name is not None else old.landlord_name,
        notes=body.notes if body.notes is not None else old.notes,
        status=_compute_status_from_dates(body.start_date, body.end_date),
    )
    db.add(new_lease)
    await db.flush()

    # Co-tenants: use explicit list if provided, otherwise carry over from old lease
    if body.co_tenant_ids is not None:
        co_ids = [cid for cid in body.co_tenant_ids if cid != new_tenant_id]
    else:
        co_ids = [ct.id for ct in old.co_tenants]
    if co_ids:
        from app.models.lease import lease_co_tenants
        from sqlalchemy import insert as sa_insert
        await db.execute(
            sa_insert(lease_co_tenants),
            [{"lease_id": new_lease.id, "user_id": cid} for cid in co_ids],
        )

    # Sync unit occupancy — mark new unit occupied, and free the old unit if it changed
    if new_lease.status == LeaseStatus.ACTIVE:
        new_unit_result = await db.execute(select(Unit).where(Unit.id == new_unit_id))
        new_unit = new_unit_result.scalar_one_or_none()
        if new_unit:
            new_unit.status = UnitStatus.OCCUPIED

    if new_unit_id != old.unit_id:
        other = await db.execute(
            select(Lease).where(
                Lease.unit_id == old.unit_id,
                Lease.id != old.id,
                Lease.status.in_([LeaseStatus.ACTIVE, LeaseStatus.PENDING]),
            )
        )
        if not other.scalar_one_or_none():
            old_unit_result = await db.execute(select(Unit).where(Unit.id == old.unit_id))
            old_unit = old_unit_result.scalar_one_or_none()
            if old_unit:
                old_unit.status = UnitStatus.VACANT

    monthly_payments = generate_monthly_payments(new_lease, member.organization_id)
    for p in monthly_payments:
        db.add(p)

    await db.commit()
    return _lease_to_out(await _get_lease(str(new_lease.id), member.organization_id, db))


# â”€â”€ Document upload / download â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

@router.post("/{lease_id}/document", response_model=LeaseOut)
async def upload_lease_document(
    lease_id: str,
    file: UploadFile = File(...),
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    lease = await _get_lease(lease_id, member.organization_id, db)

    if file.content_type not in ALLOWED_DOC_TYPES:
        raise HTTPException(status_code=400, detail="Only PDF, Word (.doc/.docx), JPEG, or PNG files accepted")

    data = await file.read()
    if len(data) > MAX_SIZE_MB * 1024 * 1024:
        raise HTTPException(status_code=400, detail=f"File exceeds {MAX_SIZE_MB} MB limit")

    # Delete old document
    if lease.document_path:
        old_path = UPLOAD_DIR / lease.document_path
        try:
            old_path.unlink(missing_ok=True)
        except OSError:
            pass

    ext = pathlib.Path(file.filename or "file").suffix or ".pdf"
    filename = f"leases/lease_{lease.id}_{uuid.uuid4()}{ext}"
    LEASE_DOC_DIR.mkdir(parents=True, exist_ok=True)
    (UPLOAD_DIR / filename).write_bytes(data)
    lease.document_path = filename

    await db.commit()
    return _lease_to_out(await _get_lease(lease_id, member.organization_id, db))


@router.delete("/{lease_id}/document", status_code=204)
async def delete_lease_document(
    lease_id: str,
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    lease = await _get_lease(lease_id, member.organization_id, db)
    if not lease.document_path:
        raise HTTPException(status_code=404, detail="No document to delete")
    try:
        (UPLOAD_DIR / lease.document_path).unlink(missing_ok=True)
    except OSError:
        pass
    lease.document_path = None
    await db.commit()



@router.post("/{lease_id}/generate-document", response_model=LeaseOut)
async def generate_lease_document(
    lease_id: str,
    body: dict,
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    """Generate a lease agreement from a template using Gemini."""
    import io
    import google.generativeai as genai
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    template_id = body.get("template_id")
    if not template_id:
        raise HTTPException(status_code=400, detail="template_id is required")

    api_key = os.environ.get("GEMINI_API_KEY", "")
    if not api_key:
        raise HTTPException(status_code=503, detail="Gemini API key not configured")

    _, member = current

    from app.models.organization import Organization as Org
    org_result = await db.execute(select(Org).where(Org.id == member.organization_id))
    org = org_result.scalar_one_or_none()
    org_landlord = org.name if org else "[LANDLORD NAME]"

    result = await db.execute(
        select(Lease)
        .where(Lease.id == lease_id, Lease.organization_id == member.organization_id)
        .options(
            selectinload(Lease.tenant),
            selectinload(Lease.co_tenants),
            selectinload(Lease.unit).selectinload(Unit.property),
        )
    )
    lease = result.scalar_one_or_none()
    if not lease:
        raise HTTPException(status_code=404, detail="Lease not found")

    landlord_name = lease.landlord_name or org_landlord

    tmpl_result = await db.execute(
        select(LeaseTemplate).where(
            LeaseTemplate.id == template_id,
            LeaseTemplate.organization_id == member.organization_id,
        )
    )
    tmpl = tmpl_result.scalar_one_or_none()
    if not tmpl:
        raise HTTPException(status_code=404, detail="Template not found")

    template_text = ""
    tmpl_path = UPLOAD_DIR / tmpl.file_path
    if tmpl_path.suffix.lower() == ".docx":
        try:
            tmpl_doc = Document(str(tmpl_path))
            template_lines = []
            for p in tmpl_doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                # Stop before the template's own signature block so it never
                # reaches the AI prompt and gets echoed back uncontrolled —
                # we always append our own deterministic SIGNATURES section below.
                lowered = text.lower()
                if lowered.startswith("landlord:") or lowered.startswith("tenant:") or lowered.startswith("tenant(s):"):
                    break
                template_lines.append(p.text)
            template_text = "\n".join(template_lines)
        except Exception:
            template_text = ""

    tenant = lease.tenant
    unit = lease.unit
    prop = unit.property if unit else None

    all_tenants = []
    if tenant:
        all_tenants.append(tenant)
    for ct in (lease.co_tenants or []):
        all_tenants.append(ct)

    tenant_names = ", ".join(
        f"{t.first_name} {t.last_name}".strip() for t in all_tenants
    ) if all_tenants else "[TENANT NAME]"

    tenant_contact_lines = []
    for i, t in enumerate(all_tenants, 1):
        label = f"Tenant {i}" if len(all_tenants) > 1 else "Tenant"
        name = f"{t.first_name} {t.last_name}".strip()
        tenant_contact_lines.append(f"{label} Name: {name}")
        tenant_contact_lines.append(f"{label} Email: {t.email or '[EMAIL]'}")
        tenant_contact_lines.append(f"{label} Tel: {t.phone or '[PHONE]'}")
    tenant_contacts = "\n".join(tenant_contact_lines) if tenant_contact_lines else "[TENANT CONTACT]"

    if prop:
        property_address = f"Unit {unit.unit_number}, {prop.address}, {prop.city}, {prop.state} {prop.zip_code}"
    elif unit:
        property_address = f"Unit {unit.unit_number}"
    else:
        property_address = "[PROPERTY ADDRESS]"

    unit_details_parts = []
    if unit:
        unit_details_parts.append(f"Unit Number: {unit.unit_number}")
        unit_details_parts.append(f"Bedrooms: {unit.bedrooms}")
        unit_details_parts.append(f"Bathrooms: {unit.bathrooms}")
        if unit.square_feet:
            unit_details_parts.append(f"Square Feet: {unit.square_feet}")
    if prop:
        unit_details_parts.append(f"Property Name: {prop.name}")
        unit_details_parts.append(f"Property Type: {prop.property_type.value if hasattr(prop.property_type, 'value') else prop.property_type}")
        unit_details_parts.append(f"Full Address: {prop.address}, {prop.city}, {prop.state} {prop.zip_code}")
    unit_details = "\n".join(unit_details_parts) if unit_details_parts else "[UNIT DETAILS]"

    lease_type_str = lease.lease_type.value if hasattr(lease.lease_type, "value") else str(lease.lease_type)

    shared_details = (
        f"TENANT NAME(S): {tenant_names}\n"
        f"TENANT CONTACT(S):\n{tenant_contacts}\n"
        f"LANDLORD NAME: {landlord_name}\n"
        f"PROPERTY DETAILS:\n{unit_details}\n"
        f"MONTHLY RENT: ${lease.monthly_rent:,.2f}\n"
        f"SECURITY DEPOSIT: ${lease.security_deposit:,.2f}\n"
        f"START DATE: {lease.start_date}\n"
        f"END DATE: {lease.end_date}\n"
        f"LEASE TYPE: {lease_type_str}\n"
    )

    instruction = (
        "IMPORTANT: The PARTIES section must list each tenant's full name, email address, and phone number exactly as provided. "
        "The PREMISES section must include the full property address, unit number, bedrooms, and bathrooms. "
        "Do not omit any of these fields."
    )

    if template_text:
        prompt = (
            "Fill in this lease agreement template with the real details below.\n\n"
            + shared_details
            + f"\n{instruction}\n\nTEMPLATE:\n\n"
            + template_text
            + "\n\nReplace all placeholders with the real values above. Keep all legal clauses. Return plain text only."
        )
    else:
        prompt = (
            "Generate a complete residential lease agreement with these details:\n\n"
            + shared_details
            + f"Style: {tmpl.description or tmpl.name}\n\n"
            + instruction
            + "\n\nUse numbered ALL CAPS section headings. Return plain text only."
        )

    try:
        genai.configure(api_key=api_key)
        model = genai.GenerativeModel("gemini-2.5-flash")
        response = model.generate_content(prompt)
        content = response.text or ""
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Gemini error: {str(e)[:200]}")

    out_doc = Document()
    title_para = out_doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("RESIDENTIAL LEASE AGREEMENT")
    run.bold = True
    run.font.size = Pt(16)
    out_doc.add_paragraph()

    # --- Parties section (always written programmatically) ---
    ph = out_doc.add_paragraph()
    ph.add_run("1. PARTIES").bold = True

    landlord_p = out_doc.add_paragraph()
    landlord_p.add_run("Landlord: ").bold = True
    landlord_p.add_run(landlord_name)

    for i, t in enumerate(all_tenants, 1):
        label = f"Tenant {i}" if len(all_tenants) > 1 else "Tenant"
        t_name = f"{t.first_name} {t.last_name}".strip()
        tp = out_doc.add_paragraph()
        tp.add_run(f"{label}: ").bold = True
        tp.add_run(t_name)
        ep = out_doc.add_paragraph(style="Normal")
        ep.paragraph_format.left_indent = Pt(18)
        ep.add_run("Email: ").bold = True
        ep.add_run(t.email or "—")
        pp = out_doc.add_paragraph(style="Normal")
        pp.paragraph_format.left_indent = Pt(18)
        pp.add_run("Tel: ").bold = True
        pp.add_run(t.phone if t.phone else "—")

    out_doc.add_paragraph()

    # --- Premises section (always written programmatically) ---
    prh = out_doc.add_paragraph()
    prh.add_run("2. PREMISES").bold = True
    addr_p = out_doc.add_paragraph()
    addr_p.add_run("Address: ").bold = True
    addr_p.add_run(property_address)
    if unit:
        beds_p = out_doc.add_paragraph()
        beds_p.add_run("Unit: ").bold = True
        beds_p.add_run(
            f"{unit.unit_number} — {unit.bedrooms} bed / {unit.bathrooms} bath"
            + (f" / {unit.square_feet} sq ft" if unit.square_feet else "")
        )
    out_doc.add_paragraph()

    # --- AI-generated body (skip sections 1 & 2 which we wrote ourselves) ---
    _skip_block = False
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            if not _skip_block:
                out_doc.add_paragraph()
            continue
        # Detect section headers
        is_header = stripped.isupper() and len(stripped) > 3
        sec_num = None
        if not is_header and len(stripped) > 2:
            parts = stripped.split(".", 1)
            if parts[0].strip().isdigit():
                is_header = True
                sec_num = int(parts[0].strip())
        if is_header:
            if sec_num in (1, 2) or "SIGNATURE" in stripped.upper():
                _skip_block = True
            else:
                _skip_block = False
                p = out_doc.add_paragraph()
                r = p.add_run(stripped)
                r.bold = True
                r.font.size = Pt(11)
        elif not _skip_block:
            out_doc.add_paragraph(stripped)

    # --- Signatures section (always written programmatically, so DocuSign can
    # reliably anchor sign/date tabs to these exact labels regardless of what
    # the AI generated) ---
    out_doc.add_paragraph()
    sh = out_doc.add_paragraph()
    sh.add_run("16. SIGNATURES").bold = True
    sh.runs[0].font.size = Pt(11)

    landlord_sig_p = out_doc.add_paragraph()
    landlord_sig_p.add_run("Landlord Signature: ________________________________    Landlord Date: ____________")

    if all_tenants:
        primary_sig_p = out_doc.add_paragraph()
        primary_sig_p.add_run("Tenant Signature: ________________________________    Tenant Date: ____________")
        for i, t in enumerate(all_tenants[1:], 2):
            co_sig_p = out_doc.add_paragraph()
            co_sig_p.add_run(f"Co-Tenant {i} Signature: ________________________________    Date: ____________")

    buf = io.BytesIO()
    out_doc.save(buf)
    doc_data = buf.getvalue()

    if lease.document_path:
        try:
            (UPLOAD_DIR / lease.document_path).unlink(missing_ok=True)
        except OSError:
            pass

    filename = f"leases/ai_lease_{lease.id}_{uuid.uuid4()}.docx"
    LEASE_DOC_DIR.mkdir(parents=True, exist_ok=True)
    (UPLOAD_DIR / filename).write_bytes(doc_data)
    lease.document_path = filename
    await db.commit()

    return _lease_to_out(await _get_lease(lease_id, member.organization_id, db))



@router.get("/{lease_id}/document")
async def download_lease_document(
    lease_id: str,
    current: tuple[User, OrganizationMember] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _, member = current
    lease = await _get_lease(lease_id, member.organization_id, db)
    if not lease.document_path:
        raise HTTPException(status_code=404, detail="No document uploaded for this lease")
    file_path = UPLOAD_DIR / lease.document_path
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Document file not found")
    return FileResponse(path=str(file_path), media_type="application/octet-stream", filename=file_path.name)


# ── DocuSign e-signature ────────────────────────────────────────────────────

@router.post("/{lease_id}/send-for-signature", response_model=LeaseOut)
async def send_lease_for_signature(
    lease_id: str,
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    from app.core import docusign

    _, member = current
    lease = await _get_lease(lease_id, member.organization_id, db)

    if not docusign.is_configured():
        raise HTTPException(status_code=503, detail="DocuSign is not configured")
    if not lease.document_path:
        raise HTTPException(status_code=400, detail="Upload or generate a lease document first")
    if not lease.tenant or not lease.tenant.email:
        raise HTTPException(status_code=400, detail="Tenant has no email on file")
    if not lease.landlord_email:
        raise HTTPException(status_code=400, detail="Set a landlord email before sending for signature")

    file_path = UPLOAD_DIR / lease.document_path
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Document file not found")
    document_bytes = file_path.read_bytes()

    landlord_name = lease.landlord_name or "Landlord"
    tenant_name = lease.tenant.display_name

    try:
        access_token = await docusign.get_access_token()
        envelope_id = await docusign.send_envelope(
            access_token,
            document_bytes,
            file_path.name,
            landlord_name,
            lease.landlord_email,
            tenant_name,
            lease.tenant.email,
            subject=f"Please sign: lease agreement for Unit {lease.unit.unit_number if lease.unit else ''}",
        )
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"DocuSign error: {str(e)[:300]}")

    lease.docusign_envelope_id = envelope_id
    lease.signature_status = "sent"
    await db.commit()

    return _lease_to_out(await _get_lease(lease_id, member.organization_id, db))


@router.post("/{lease_id}/signature-status", response_model=LeaseOut)
async def check_lease_signature_status(
    lease_id: str,
    current: tuple[User, OrganizationMember] = Depends(require_min_role(UserRole.OWNER)),
    db: AsyncSession = Depends(get_db),
):
    from app.core import docusign

    _, member = current
    lease = await _get_lease(lease_id, member.organization_id, db)

    if not docusign.is_configured():
        raise HTTPException(status_code=503, detail="DocuSign is not configured")
    if not lease.docusign_envelope_id:
        raise HTTPException(status_code=400, detail="This lease has not been sent for signature yet")

    try:
        access_token = await docusign.get_access_token()
        envelope_status = await docusign.get_envelope_status(access_token, lease.docusign_envelope_id)

        resolved_status = envelope_status
        if envelope_status not in ("completed", "declined", "voided"):
            # Distinguish "tenant signed, awaiting landlord" from the envelope's
            # overall status, which stays "sent"/"delivered" until everyone signs.
            signers = await docusign.get_signer_statuses(access_token, lease.docusign_envelope_id)
            tenant_signer = next((s for s in signers if s.get("recipientId") == "2"), None)
            if tenant_signer and tenant_signer.get("status") == "completed":
                resolved_status = "tenant_signed"

        if envelope_status == "completed" and lease.signature_status != "completed":
            # DocuSign's combined-document download is always a PDF regardless of the
            # original file type, so it gets its own .pdf path rather than overwriting
            # the old file's bytes under whatever extension it originally had.
            signed_bytes = await docusign.get_combined_document(access_token, lease.docusign_envelope_id)
            old_path = lease.document_path
            new_filename = f"leases/signed_lease_{lease.id}_{uuid.uuid4()}.pdf"
            LEASE_DOC_DIR.mkdir(parents=True, exist_ok=True)
            (UPLOAD_DIR / new_filename).write_bytes(signed_bytes)
            lease.document_path = new_filename
            if old_path:
                try:
                    (UPLOAD_DIR / old_path).unlink(missing_ok=True)
                except OSError:
                    pass
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"DocuSign error: {str(e)[:300]}")

    lease.signature_status = resolved_status
    await db.commit()

    return _lease_to_out(await _get_lease(lease_id, member.organization_id, db))
